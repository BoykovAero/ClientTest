"""Извлечение текста из присланных файлов.

Word и Excel разбираются локально, без обращения к моделям: так точнее и
бесплатно. Полученный текст дальше идёт через тот же разбор, что и обычное
сообщение.
"""

from __future__ import annotations

import csv
import io
import logging
from pathlib import PurePosixPath

from bot.grid import fill_merges, find_column, narrow, to_text

logger = logging.getLogger(__name__)

# Telegram отдаёт ботам файлы не больше 20 МБ, но расписание столько не весит.
MAX_FILE_BYTES = 10 * 1024 * 1024
# Предел на объём текста. Он высокий, потому что разбор идёт частями:
# в модель за раз уходит кусок, а не весь файл. Ограничение остаётся ради
# страховки от по-настоящему гигантских выгрузок.
MAX_TEXT_CHARS = 60_000

DOCUMENT_SUFFIXES = {".docx", ".xlsx", ".txt", ".md", ".csv"}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".heic"}


class DocumentError(RuntimeError):
    """Файл не удалось прочитать. Текст уходит пользователю."""


def kind_of(filename: str, mime_type: str = "") -> str:
    """Что это за файл: 'document', 'image' или 'unsupported'."""
    suffix = PurePosixPath(filename or "").suffix.lower()
    if suffix in DOCUMENT_SUFFIXES:
        return "document"
    if suffix in IMAGE_SUFFIXES or (mime_type or "").startswith("image/"):
        return "image"
    return "unsupported"


def extract_text(filename: str, data: bytes, instruction: str = "") -> tuple[str, str]:
    """Файл -> (текст, имя выбранного столбца).

    Если в просьбе назван столбец таблицы, в текст попадает только он и
    столбец времени: расписание класса — это одна колонка из двадцати.
    """
    if not data:
        raise DocumentError("файл пустой")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError(f"файл больше {MAX_FILE_BYTES // (1024 * 1024)} МБ")

    column_name = ""
    suffix = PurePosixPath(filename or "").suffix.lower()
    if suffix == ".docx":
        text = _from_docx(data)
    elif suffix == ".xlsx":
        text, column_name = _from_xlsx(data, instruction)
    elif suffix in {".txt", ".md"}:
        text = _decode(data)
    elif suffix == ".csv":
        text = _from_csv(data)
    else:
        raise DocumentError(
            f"формат {suffix or '?'} не поддерживается. "
            "Пришли .docx, .xlsx, .csv, .txt или картинку"
        )

    text = text.strip()
    if not text:
        raise DocumentError(
            "в файле не нашлось текста. Если это скан или картинка внутри "
            "документа — пришли её отдельно как изображение"
        )

    if len(text) > MAX_TEXT_CHARS:
        logger.info("Файл %r обрезан с %d до %d символов", filename, len(text), MAX_TEXT_CHARS)
        text = text[:MAX_TEXT_CHARS]
    return text, column_name


def is_large(text: str) -> bool:
    """Стоит ли предупредить пользователя, что разбор займёт время."""
    from bot.parser import CHUNK_CHARS

    return len(text) > CHUNK_CHARS


# Доля управляющих символов, выше которой считаем, что перед нами не текст.
MAX_CONTROL_RATIO = 0.05


def _looks_like_text(text: str) -> bool:
    """Отсекает двоичные файлы, притворившиеся текстом.

    cp1251 однобайтовая и «расшифровывает» практически любой мусор, поэтому
    успешного decode мало: без этой проверки картинка, названная .txt, ушла
    бы в модель кракозябрами.
    """
    if not text:
        return False
    control = sum(1 for ch in text if ch < " " and ch not in "\n\r\t")
    return control / len(text) <= MAX_CONTROL_RATIO


def _decode(data: bytes) -> str:
    """Текст в неизвестной кодировке. Windows-1251 всё ещё встречается."""
    for encoding in ("utf-8", "utf-8-sig", "cp1251"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if _looks_like_text(text):
            return text
    raise DocumentError("файл не похож на текст — проверь формат и кодировку")


def _from_csv(data: bytes) -> str:
    rows = csv.reader(io.StringIO(_decode(data)))
    return "\n".join(" | ".join(cell.strip() for cell in row if cell.strip()) for row in rows)


def _from_docx(data: bytes) -> str:
    """Абзацы и таблицы. Расписания часто лежат именно в таблицах."""
    try:
        import docx
    except ImportError:
        raise DocumentError("на сервере не установлен python-docx") from None

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(f"файл не читается как .docx ({exc})") from exc

    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _from_xlsx(data: bytes, instruction: str = "") -> tuple[str, str]:
    """Листы книги. Если в просьбе назван столбец — только он.

    Читается без read_only: объединённые области нужны целиком, а в
    режиме чтения openpyxl их не отдаёт.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise DocumentError("на сервере не установлен openpyxl") from None

    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise DocumentError(f"файл не читается как .xlsx ({exc})") from exc

    parts: list[str] = []
    column_name = ""
    try:
        for sheet in workbook.worksheets:
            rows = [
                [" ".join(str(cell.value).split()) if cell.value is not None else ""
                 for cell in row]
                for row in sheet.iter_rows()
            ]
            if not rows:
                continue
            merges = [
                (m.min_row - 1, m.max_row - 1, m.min_col - 1, m.max_col - 1)
                for m in sheet.merged_cells.ranges
            ]
            grid = fill_merges(rows, merges)

            found = find_column(grid, instruction) if instruction else None
            if found is not None:
                column, column_name = found
                body = narrow(grid, column)
            else:
                body = to_text(grid)

            if not body:
                continue
            if len(workbook.worksheets) > 1:
                parts.append(f"# лист: {sheet.title}")
            parts.append(body)
    finally:
        workbook.close()
    return "\n".join(parts), column_name
